import unittest
from datetime import date
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from tools.smt_model_comparison import build_comparison_workbook, build_dimension_rows, defect_ppm


class SMTModelComparisonTests(unittest.TestCase):
    def test_categories_align_and_count_unique_pcbs_per_model(self):
        views = {
            "A": {"confirmed": pd.DataFrame({
                "_DefectKey": ["a1", "a1", "a2"],
                "Operation": ["Download", "Download", "AOI"],
            })},
            "B": {"confirmed": pd.DataFrame({
                "_DefectKey": ["b1", "b2"],
                "Operation": ["Download", None],
            })},
        }
        rows = build_dimension_rows(views, "Operation")
        self.assertEqual(rows[0], {
            "category": "Download", "counts": {"A": 1, "B": 1}, "total": 2,
        })
        self.assertEqual(next(row for row in rows if row["category"] == "AOI")["counts"], {"A": 1, "B": 0})
        self.assertEqual(next(row for row in rows if row["category"] == "Not specified")["counts"], {"A": 0, "B": 1})

    def test_ppm_requires_valid_input(self):
        self.assertEqual(defect_ppm(3, 200), 15_000)
        self.assertIsNone(defect_ppm(1, 0))
        self.assertIsNone(defect_ppm(3, 2))

    def test_formatted_workbook_preserves_numeric_values_and_targets(self):
        metrics = [
            {"label": "Functional Pass Rate", "kind": "rate", "target": 0.9966,
             "lower_is_better": False, "values": {"A": 0.98, "B": 1.0}},
            {"label": "Defect PCB", "kind": "count", "target": None,
             "lower_is_better": False, "values": {"A": 2, "B": 0}},
        ]
        dimensions = [("TestOperation", [
            {"category": "Download", "counts": {"A": 2, "B": 0}, "total": 2},
        ])]
        payload = build_comparison_workbook(
            ["A", "B"], date(2026, 9, 1), date(2026, 9, 9), metrics, dimensions,
            {"A": 200, "B": 100},
        )
        workbook = load_workbook(BytesIO(payload))
        self.assertEqual(workbook.sheetnames, ["KPI Summary", "Defect Breakdown"])
        summary = workbook["KPI Summary"]
        self.assertEqual(summary["C7"].value, 0.98)
        self.assertEqual(summary["D7"].value, 1.0)
        self.assertEqual(summary["C7"].number_format, "0.00%")
        self.assertNotEqual(summary["C7"].fill.fgColor.rgb, summary["D7"].fill.fgColor.rgb)
        self.assertEqual(summary["C8"].value, 2)
        detail = workbook["Defect Breakdown"]
        self.assertEqual(detail["B7"].value, "Download")
        self.assertEqual(detail["C7"].value, 2)
        self.assertEqual(detail["D7"].value, 10_000)
        self.assertEqual(detail["E7"].value, 0)
        self.assertEqual(detail["F7"].value, 0)
        self.assertEqual(detail.auto_filter.ref, "A6:F7")

    def test_workbook_title_uses_requested_area(self):
        payload = build_comparison_workbook(
            ["A", "B"], date(2026, 9, 1), date(2026, 9, 9), [], [],
            {"A": 100, "B": 100}, area="Assembly",
        )
        workbook = load_workbook(BytesIO(payload))
        self.assertEqual(workbook["KPI Summary"]["A1"].value, "JOVI · Assembly Model Comparison")

    def test_combined_report_uses_each_areas_input_for_ppm(self):
        dimensions = [
            ("SMT · TestOperation", [{"category": "AOI", "counts": {"A": 2, "B": 1}}], {"A": 100, "B": 100}),
            ("Assembly · TestOperation", [{"category": "Audio", "counts": {"A": 2, "B": 1}}], {"A": 200, "B": 500}),
        ]
        payload = build_comparison_workbook(
            ["A", "B"], date(2026, 9, 1), date(2026, 9, 9), [], dimensions,
            {"A": 0, "B": 0}, area="SMT + Assembly",
        )
        breakdown = load_workbook(BytesIO(payload))["Defect Breakdown"]
        self.assertEqual(breakdown["D7"].value, 20_000)
        self.assertEqual(breakdown["D8"].value, 10_000)
        self.assertEqual(breakdown["F8"].value, 2_000)


if __name__ == "__main__":
    unittest.main()
