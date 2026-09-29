"""Side-by-side model comparison using each area's confirmed defect records."""

from __future__ import annotations

from math import ceil

import pandas as pd


def build_dimension_rows(
    views: dict[str, dict], column: str
) -> list[dict[str, object]]:
    """Count unique NG PCBs per category and model, preserving zeroes for missing categories."""
    counts_by_model: dict[str, dict[str, int]] = {}
    categories: set[str] = set()
    for model, view in views.items():
        confirmed = view["confirmed"]
        if confirmed.empty:
            counts_by_model[model] = {}
            continue
        values = confirmed.get(column, pd.Series("", index=confirmed.index))
        normalized = values.fillna("").astype(str).str.strip()
        normalized = normalized.mask(normalized.eq("") | normalized.str.lower().isin({"nan", "none"}), "Not specified")
        grouped = (
            confirmed.assign(_Category=normalized)
            .groupby("_Category")["_DefectKey"]
            .nunique()
        )
        counts_by_model[model] = {str(category): int(count) for category, count in grouped.items()}
        categories.update(counts_by_model[model])

    rows = []
    for category in categories:
        model_counts = {model: counts_by_model[model].get(category, 0) for model in views}
        rows.append({"category": category, "counts": model_counts, "total": sum(model_counts.values())})
    return sorted(rows, key=lambda row: (-row["total"], row["category"].casefold()))


def defect_ppm(count: int, input_count: int) -> float | None:
    return count / input_count * 1_000_000 if input_count > 0 and count <= input_count else None


def build_comparison_workbook(
    models: list[str],
    start_date,
    end_date,
    metrics: list[dict],
    dimensions: list[tuple],
    inputs: dict[str, int],
    area: str = "SMT",
) -> bytes:
    """Create a formatted, printable Excel snapshot of the selected comparison."""
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    navy, blue, pale_blue = "0B2D5A", "246DCF", "EAF2FF"
    text, muted, border = "17375F", "607A9F", "D6E2F1"
    green, pale_green, red, pale_red = "087C4B", "E8F8EF", "B91C1C", "FDECEC"
    white, stripe = "FFFFFF", "F7FAFE"
    thin_bottom = Border(bottom=Side(style="hair", color=border))
    workbook = Workbook()
    summary = workbook.active
    summary.title = "KPI Summary"
    breakdown = workbook.create_sheet("Defect Breakdown")
    last_summary_col = get_column_letter(len(models) + 2)
    last_breakdown_col = get_column_letter(2 + 2 * len(models))
    period_text = f"Analysis period: {start_date:%d/%m/%Y} – {end_date:%d/%m/%Y}"

    def title_band(sheet, last_col: str, title: str) -> None:
        for row in sheet.iter_rows(min_row=1, max_row=2, min_col=1, max_col=sheet[last_col + "1"].column):
            for band_cell in row:
                band_cell.fill = PatternFill("solid", fgColor=navy)
        sheet.merge_cells(f"A1:{last_col}2")
        cell = sheet["A1"]
        cell.value = title
        cell.font = Font(name="Aptos Display", size=20, bold=True, color=white)
        cell.alignment = Alignment(vertical="center", indent=1)
        sheet.row_dimensions[1].height = 24
        sheet.row_dimensions[2].height = 24
        sheet.merge_cells(f"A3:{last_col}3")
        sheet["A3"] = period_text
        sheet["A3"].font = Font(name="Aptos", size=11, bold=True, color=blue)
        sheet["A3"].alignment = Alignment(vertical="center")
        sheet.row_dimensions[3].height = 22
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_options.horizontalCentered = True
        sheet.page_margins.left = sheet.page_margins.right = 0.3

    def header_row(sheet, row_number: int, headers: list[str]) -> None:
        for column_number, label in enumerate(headers, 1):
            cell = sheet.cell(row_number, column_number, label)
            cell.fill = PatternFill("solid", fgColor=blue)
            cell.font = Font(name="Aptos", size=10, bold=True, color=white)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[row_number].height = 33

    def body_cell(cell, row_number: int, *, strong: bool = False) -> None:
        cell.fill = PatternFill("solid", fgColor=white if row_number % 2 else stripe)
        cell.font = Font(name="Aptos", size=10, bold=strong, color=text)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = thin_bottom

    title_band(summary, last_summary_col, f"JOVI · {area} Model Comparison")
    summary.merge_cells(f"A4:{last_summary_col}4")
    denominator_note = (
        "rates use each area's own model input" if area == "SMT + Assembly"
        else f"rates use each model's own {area} input"
    )
    summary["A4"] = f"Same period and MES rules for every model · {denominator_note}"
    summary["A4"].font = Font(name="Aptos", size=9, italic=True, color=muted)
    header_row(summary, 6, ["KPI", "Target", *models])
    summary.column_dimensions["A"].width = 31
    summary.column_dimensions["B"].width = 17
    for column_number in range(3, len(models) + 3):
        summary.column_dimensions[get_column_letter(column_number)].width = 30

    for row_number, metric in enumerate(metrics, 7):
        kind = metric["kind"]
        label_cell = summary.cell(row_number, 1, metric["label"])
        target_cell = summary.cell(row_number, 2, metric.get("target"))
        body_cell(label_cell, row_number, strong=True)
        body_cell(target_cell, row_number)
        number_format = "0.00%" if kind == "rate" else '#,##0" PPM"' if kind == "ppm" else "#,##0"
        if target_cell.value is not None:
            target_cell.number_format = number_format
        for column_number, model in enumerate(models, 3):
            value = metric["values"].get(model)
            cell = summary.cell(row_number, column_number, value if value is not None else "N/A")
            body_cell(cell, row_number, strong=True)
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                cell.number_format = number_format
                target = metric.get("target")
                if target is not None:
                    good = value <= target if metric.get("lower_is_better") else value >= target
                    cell.fill = PatternFill("solid", fgColor=pale_green if good else pale_red)
                    cell.font = Font(name="Aptos", size=10, bold=True, color=green if good else red)
            else:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        summary.row_dimensions[row_number].height = 58 if kind == "text" else 27
    note_row = 8 + len(metrics)
    summary.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=len(models) + 2)
    summary.cell(note_row, 1, "Green = on target · red = off target · N/A = unavailable input or invalid rate")
    summary.cell(note_row, 1).font = Font(name="Aptos", size=9, italic=True, color=muted)
    summary.freeze_panes = "C7"
    summary.print_title_rows = "1:6"
    summary.print_area = f"A1:{last_summary_col}{note_row}"

    title_band(breakdown, last_breakdown_col, f"JOVI · {area} Defect Breakdown")
    breakdown.merge_cells(f"A4:{last_breakdown_col}4")
    breakdown["A4"] = "NG PCBs are unique within each category; one PCB may appear in multiple categories. PPM uses each model's input."
    breakdown["A4"].font = Font(name="Aptos", size=9, italic=True, color=muted)
    headers = ["Dimension", "Defect category"]
    for model in models:
        headers.extend([f"{model}\nNG PCBs", f"{model}\nPPM"])
    header_row(breakdown, 6, headers)
    breakdown.column_dimensions["A"].width = 26
    breakdown.column_dimensions["B"].width = 52
    for column_number in range(3, 2 * len(models) + 3):
        breakdown.column_dimensions[get_column_letter(column_number)].width = 18
    row_number = 7
    for entry in dimensions:
        dimension, rows = entry[:2]
        dimension_inputs = entry[2] if len(entry) > 2 else inputs
        for row in rows:
            values = [dimension, row["category"]]
            for model in models:
                count = row["counts"][model]
                values.extend([count, defect_ppm(count, dimension_inputs.get(model, 0))])
            for column_number, value in enumerate(values, 1):
                cell = breakdown.cell(row_number, column_number, value if value is not None else "N/A")
                body_cell(cell, row_number, strong=column_number <= 2)
                if column_number > 2:
                    cell.number_format = '#,##0" PPM"' if column_number % 2 == 0 else "#,##0"
            breakdown.row_dimensions[row_number].height = max(30, min(90, 15 * ceil(len(str(row["category"])) / 48)))
            row_number += 1
    breakdown.auto_filter.ref = f"A6:{last_breakdown_col}{max(6, row_number - 1)}"
    breakdown.freeze_panes = "C7"
    breakdown.print_title_rows = "1:6"
    breakdown.print_area = f"A1:{last_breakdown_col}{max(6, row_number - 1)}"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
